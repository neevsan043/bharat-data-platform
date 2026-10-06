"""Landing-zone (file drop) ingesters: NPCI monthly Excel and TRAI PDF reports.

New files are detected by sha256 against a per-source ledger, the layout is
validated loudly before anything is written, bronze partitions are written
atomically, and processed files are moved to data/landing/_processed/<source>/.
"""

from __future__ import annotations

import json
import logging
import os
import re
import uuid
from abc import abstractmethod
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any, ClassVar

import pyarrow as pa

from ingestion.base import (
    Ingester,
    IngestionError,
    RunResult,
    SchemaChangedError,
    sha256_file,
)
from ingestion.report_tables import REPORT_ROWS_COLUMNS, extract_pdf_rows

LOGGER = logging.getLogger(__name__)

_MONTHS = {
    "January": 1, "February": 2, "March": 3, "April": 4, "May": 5, "June": 6,
    "July": 7, "August": 8, "September": 9, "October": 10, "November": 11, "December": 12,
}
_MONTH_RE = re.compile(r"^([A-Z][a-z]+)-(\d{4})$")
_FY_RE = re.compile(r"(20\d{2}-\d{2})")


def file_sha256(path: Path) -> str:
    return sha256_file(path)


def parse_month(value: str) -> date:
    """'March-2022' -> date(2022, 3, 1). Anything else is a schema break."""
    match = _MONTH_RE.match(value.strip())
    if not match or match.group(1) not in _MONTHS:
        raise SchemaChangedError(f"npci: unexpected Month value {value!r} (expected e.g. 'March-2022')")
    return date(int(match.group(2)), _MONTHS[match.group(1)], 1)


def to_float(value: object, column: str) -> float:
    """Xlsx floats pass through; strings with Indian separators are tolerated once."""
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError as exc:
        raise SchemaChangedError(f"npci: non-numeric {column} value {value!r}") from exc


class LandingIngester(Ingester):
    """Common file-drop flow: checksum ledger -> ingest -> move to _processed."""

    landing_glob: ClassVar[str] = "*"

    def plan(self, **_: Any) -> list[Any]:
        # Landing runs are checksum-driven (see run()), not partition-planned.
        return []

    # ------------------------------------------------------------- ledger
    def _ledger_path(self) -> Path:
        return self.settings.manifest_dir / f"landing_{self.source}.jsonl"

    def _rejects_path(self) -> Path:
        return self.settings.manifest_dir / f"landing_{self.source}_rejects.jsonl"

    def _processed_checksums(self) -> set[str]:
        ledger = self._ledger_path()
        if not ledger.exists():
            return set()
        with ledger.open(encoding="utf-8") as handle:
            return {json.loads(line)["sha256"] for line in handle if line.strip()}

    def _mark_processed(self, path: Path, digest: str, partitions: list[str]) -> None:
        entry = {
            "filename": path.name,
            "sha256": digest,
            "processed_at": datetime.now(UTC).isoformat(timespec="seconds"),
            "partitions": partitions,
        }
        ledger = self._ledger_path()
        ledger.parent.mkdir(parents=True, exist_ok=True)
        with ledger.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(entry, sort_keys=True) + "\n")

    def _mark_rejected(self, path: Path, digest: str, reason: str) -> None:
        entry = {
            "filename": path.name,
            "sha256": digest,
            "rejected_at": datetime.now(UTC).isoformat(timespec="seconds"),
            "reason": reason,
        }
        rejects = self._rejects_path()
        rejects.parent.mkdir(parents=True, exist_ok=True)
        with rejects.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(entry, sort_keys=True) + "\n")

    # ------------------------------------------------------------- detection
    def pending_files(self) -> list[tuple[Path, str]]:
        landing = self.settings.landing_source_dir(self.source)
        if not landing.exists():
            return []
        processed = self._processed_checksums()
        pending = []
        for path in sorted(landing.glob(self.landing_glob)):
            if not path.is_file():
                continue
            digest = file_sha256(path)
            if digest not in processed:
                pending.append((path, digest))
        return pending

    # ------------------------------------------------------------- runner
    def run(self, plans: Any = None, *, overwrite: bool = False, only: Any = None) -> list[RunResult]:
        _ = plans, only, overwrite  # landing runs are checksum-driven
        run_id = uuid.uuid4().hex[:12]
        results: list[RunResult] = []
        pending = self.pending_files()
        if not pending:
            LOGGER.info("%s: no new files in data/landing/%s/", self.source, self.source)
            return results
        for path, digest in pending:
            LOGGER.info("%s: processing %s (sha256 %s…)", self.source, path.name, digest[:12])
            try:
                written = self.ingest_file(path, digest, run_id)
            except SchemaChangedError:
                raise  # a layout change is a loud failure, not a quarantine case
            except IngestionError as exc:
                # One unparseable file must not block the rest of the landing set:
                # record the rejection, quarantine the file, keep going.
                LOGGER.error("%s: REJECTED %s: %s", self.source, path.name, exc)
                rejected_dir = self.settings.landing_processed_dir(self.source) / "_rejected"
                rejected_dir.mkdir(parents=True, exist_ok=True)
                destination = rejected_dir / path.name
                if destination.exists():
                    raise IngestionError(
                        f"{self.source}: _rejected already contains {destination.name} - resolve manually"
                    ) from exc
                os.replace(path, destination)
                self._mark_rejected(destination, digest, str(exc))
                continue
            processed_dir = self.settings.landing_processed_dir(self.source)
            processed_dir.mkdir(parents=True, exist_ok=True)
            destination = processed_dir / path.name
            if destination.exists():
                raise IngestionError(
                    f"{self.source}: _processed already contains {destination.name} "
                    "(same name, different content?) - resolve manually"
                )
            os.replace(path, destination)
            self._mark_processed(destination, digest, [r.partition for r in written])
            LOGGER.info("%s: moved %s -> %s", self.source, path.name, destination)
            results.extend(written)
        return results

    @abstractmethod
    def ingest_file(self, path: Path, digest: str, run_id: str) -> list[RunResult]:
        """Validate and write one landing file; raise loudly on any layout change."""


class NpciIngester(LandingIngester):
    """NPCI UPI monthly statistics Excel (real files in data/landing/npci/)."""

    source = "npci"
    landing_glob = "*.xlsx"

    EXPECTED_HEADERS: ClassVar[tuple[str, ...]] = (
        "Month",
        "Volume (In Mn.)",
        "Avg. Daily Volume (In Mn.)",
        "Value (In Cr.)",
        "Avg. Daily Value (In Cr.)",
    )

    SCHEMA = pa.schema(
        [
            ("month_start", pa.date32()),
            ("fy", pa.string()),
            ("volume_mn", pa.float64()),
            ("avg_daily_volume_mn", pa.float64()),
            ("value_cr", pa.float64()),
            ("avg_daily_value_cr", pa.float64()),
            ("source_file", pa.string()),
        ]
    )

    def ingest_file(self, path: Path, digest: str, run_id: str) -> list[RunResult]:
        from openpyxl import load_workbook

        fy_match = _FY_RE.search(path.name)
        if not fy_match:
            raise SchemaChangedError(f"npci: filename {path.name!r} has no FY tag (20XX-YY)")
        fy = fy_match.group(1)

        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        assert worksheet is not None
        rows_iter = worksheet.iter_rows(values_only=True)
        headers = next(rows_iter, None)
        if headers is None:
            raise SchemaChangedError(f"npci: {path.name} is empty")
        actual = tuple(str(h).strip() if h is not None else "" for h in headers[: len(self.EXPECTED_HEADERS)])
        if actual != self.EXPECTED_HEADERS:
            raise SchemaChangedError(
                f"npci: layout changed in {path.name}\n"
                f"  expected: {self.EXPECTED_HEADERS}\n"
                f"  found:    {actual}"
            )
        records: list[dict[str, Any]] = []
        for row in rows_iter:
            if row is None or all(cell is None for cell in row):
                continue
            month_raw = row[0]
            if month_raw is None:
                raise SchemaChangedError(f"npci: blank Month cell in {path.name}")
            records.append(
                {
                    "month_start": parse_month(str(month_raw)),
                    "fy": fy,
                    "volume_mn": to_float(row[1], "Volume"),
                    "avg_daily_volume_mn": to_float(row[2], "Avg. Daily Volume"),
                    "value_cr": to_float(row[3], "Value"),
                    "avg_daily_value_cr": to_float(row[4], "Avg. Daily Value"),
                    "source_file": path.name,
                }
            )
        workbook.close()
        table = pa.Table.from_pylist(records, schema=self.SCHEMA)
        bad = [r for r in records if r["volume_mn"] <= 0 or r["value_cr"] <= 0]
        if bad:
            raise IngestionError(f"npci: non-positive volumes/values in {path.name}: {bad[:3]}")
        partition = f"fy={fy}"
        out_path = self.partition_path(partition)
        fetched_at = datetime.now(UTC).isoformat(timespec="seconds")
        checksum = self.atomic_write_table(out_path, table)
        result = RunResult(
            self.source, partition, str(out_path), table.num_rows, checksum, fetched_at, run_id, "written"
        )
        self.append_manifest(result)
        LOGGER.info("npci/%s: wrote %d rows from %s", partition, table.num_rows, path.name)
        return [result]


class TraiPdfIngester(LandingIngester):
    """TRAI performance/subscription PDFs dropped into data/landing/trai/."""

    source = "trai"
    landing_glob = "*.pdf"
    required_columns = REPORT_ROWS_COLUMNS

    def ingest_file(self, path: Path, digest: str, run_id: str) -> list[RunResult]:
        title = path.stem
        table = extract_pdf_rows(
            path,
            report_type="trai_pdf",
            report_title=title,
            report_date=date.today(),  # report date is metadata, refined in silver
            source_label=path.name,
            max_pages=250,  # QPIRs run to 214 pages
        )
        partition = f"drop={digest[:10]}"
        out_path = self.partition_path(partition)
        fetched_at = datetime.now(UTC).isoformat(timespec="seconds")
        checksum = self.atomic_write_table(out_path, table)
        result = RunResult(
            self.source, partition, str(out_path), table.num_rows, checksum, fetched_at, run_id, "written"
        )
        self.append_manifest(result)
        LOGGER.info("trai/%s: wrote %d rows from %s", partition, table.num_rows, path.name)
        return [result]
