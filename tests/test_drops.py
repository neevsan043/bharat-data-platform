"""Unit tests for the landing-zone ingesters - real NPCI workbook, no network."""

import shutil
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from ingestion.base import IngestionError, SchemaChangedError
from ingestion.config import Settings
from ingestion.drops import NpciIngester, parse_month, to_float

REPO_ROOT = Path(__file__).resolve().parents[1]
SAMPLE = REPO_ROOT / "data" / "samples" / "npci_monthly_2021-22_sample.xlsx"


def _settings(tmp_path: Path) -> Settings:
    return Settings(
        repo_root=REPO_ROOT,
        data_dir=tmp_path,
        user_agent="test-agent",
        request_timeout_s=30,
        opencellid_token=None,
    )


def test_parse_month() -> None:
    assert parse_month("March-2022") == date(2022, 3, 1)
    assert parse_month(" April-2026 ") == date(2026, 4, 1)
    with pytest.raises(SchemaChangedError):
        parse_month("2022-03")


def test_to_float() -> None:
    assert to_float(5405.65, "Volume") == 5405.65
    assert to_float("29,37,396.69", "Value") == 2937396.69
    with pytest.raises(SchemaChangedError):
        to_float("n/a", "Value")


@pytest.mark.skipif(not SAMPLE.exists(), reason="sample not present (git-ignored data/)")
def test_ingest_real_landing_file_and_idempotency(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    landing = settings.landing_source_dir("npci")
    landing.mkdir(parents=True)
    shutil.copy(SAMPLE, landing / SAMPLE.name)

    ingester = NpciIngester(settings)
    # The 2021-22 file holds 10 months (June-2021 .. March-2022), verified in Phase 0.
    results = ingester.run()
    assert len(results) == 1
    assert results[0].partition == "fy=2021-22"
    assert results[0].status == "written"
    assert results[0].row_count == 10

    bronze_file = settings.bronze_source_dir("npci") / "fy=2021-22.parquet"
    assert bronze_file.exists()

    # Re-running must not duplicate: the same checksum is already in the ledger.
    second = ingester.run()
    assert second == []

    # The processed file has moved out of landing into _processed.
    assert not (landing / SAMPLE.name).exists()
    assert (settings.landing_processed_dir("npci") / SAMPLE.name).exists()


@pytest.mark.skipif(not SAMPLE.exists(), reason="sample not present (git-ignored data/)")
def test_layout_change_fails_loudly(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    landing = settings.landing_source_dir("npci")
    landing.mkdir(parents=True)
    broken = landing / "Product-Statistics-UPI-Upi-monthly-statistics-2099-00-monthly.xlsx"

    source_wb = load_workbook(SAMPLE, read_only=True)
    target_wb = Workbook()
    target_ws = target_wb.active
    assert target_ws is not None
    source_ws = source_wb.active
    assert source_ws is not None
    for row_index, row in enumerate(source_ws.iter_rows(values_only=True), start=1):
        for col_index, value in enumerate(row, start=1):
            target_ws.cell(row=row_index, column=col_index, value=value)
    # simulate an upstream layout change: rename a column
    target_ws.cell(row=1, column=2, value="Vol (Mn)")
    target_wb.save(broken)

    with pytest.raises(SchemaChangedError, match="layout changed"):
        NpciIngester(settings).run()


def test_non_positive_values_quarantined(tmp_path: Path) -> None:
    """A data defect (not a layout change) is quarantined to _rejected, not ingested."""
    settings = _settings(tmp_path)
    landing = settings.landing_source_dir("npci")
    landing.mkdir(parents=True)
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(list(NpciIngester.EXPECTED_HEADERS))
    worksheet.append(["March-2099", -5.0, 1.0, 1.0, 1.0])
    workbook.save(landing / "Product-Statistics-UPI-Upi-monthly-statistics-2099-00-monthly.xlsx")
    results = NpciIngester(settings).run()
    assert results == []
    rejected = settings.landing_processed_dir("npci") / "_rejected"
    assert (rejected / "Product-Statistics-UPI-Upi-monthly-statistics-2099-00-monthly.xlsx").exists()
    ledger = settings.manifest_dir / "landing_npc i_rejects.jsonl"
    ledger = settings.manifest_dir / "landing_npci_rejects.jsonl"
    assert "non-positive" in ledger.read_text()
    _ = ledger if ledger.exists() else IngestionError
